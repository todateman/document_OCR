"""PDF・Excel の入力（document_input.py）のテスト。

電子 PDF とスキャン PDF は電子天秤のサンプル画像から作り、Excel は openpyxl で様式を作る。
"""
import copy
import sys
from pathlib import Path

import numpy as np
import pytest
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

import document_input as di  # noqa: E402
import inspection_ocr as io  # noqa: E402
import template_ocr as t  # noqa: E402
from ocr_common import CHECK, NG, OK, load_config  # noqa: E402
from test_template import CONFIG, EXPECTED, SAMPLE, items  # noqa: E402

NUMBERS = {k: v for k, v in EXPECTED.items() if v[-1:].isdigit()}  # 目視で確認した数値の項目


# ---------------------------------------------------------------- PDF

@pytest.fixture(scope="module")
def cfg():
    return load_config(CONFIG)


@pytest.fixture(scope="module")
def aligned(cfg):
    return t.align(t.load_page(SAMPLE), cfg)


@pytest.fixture(scope="module")
def truth(cfg, aligned):
    """全項目の枠と、元画像から読んだ値（テストで目視の値と一致を確認済み）。"""
    ref, _, _ = t.load_reference(cfg)
    rep = t.evaluate(aligned, cfg, "rapidocr")
    read = {it.name: it.cell.text for it in rep.items}
    return [(tg.name, tg.box, read[tg.name]) for _, targets in t.expand(cfg, ref) for tg in targets]


def _pdf(path, gray, texts=(), *, color=0, invisible=False, whiten=False):
    """gray を 1 px = 1 pt で背景に置き、texts [(枠, 文字列)] を文字として書き込んだ PDF。"""
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfgen import canvas

    pdfmetrics.registerFont(UnicodeCIDFont("HeiseiKakuGo-W5"))
    bg = gray.copy()
    if whiten:  # 値の欄を消して、文字だけで値を表す（罫線は残す）
        for (x0, y0, x1, y1), _ in texts:
            bg[y0 + 4 : y1 - 4, x0 + 4 : x1 - 4] = 255
    h, w = bg.shape
    png = path.with_suffix(".png")
    Image.fromarray(bg).save(png)
    c = canvas.Canvas(str(path), pagesize=(w, h))
    c.drawImage(ImageReader(str(png)), 0, 0, w, h)
    for (x0, y0, x1, y1), text in texts:
        size = min((y1 - y0) * 0.55, 22)
        to = c.beginText()
        to.setTextRenderMode(3 if invisible else 0)
        to.setFont("HeiseiKakuGo-W5", size)
        to.setFillGray(color)
        tw = c.stringWidth(text, "HeiseiKakuGo-W5", size)
        to.setTextOrigin((x0 + x1 - tw) / 2, h - (y0 + y1) / 2 - size * 0.35)
        to.textLine(text)
        c.drawText(to)
    c.save()
    return path


def test_digital_pdf_reads_text_layer(cfg, aligned, truth, tmp_path):
    # 値を白い文字で書く（画面では見えないので OCR では読めない）→ 読めればテキスト層を使っている
    texts = [(box, v) for _, box, v in truth if v]
    pdf = _pdf(tmp_path / "digital.pdf", aligned, texts, color=1, whiten=True)
    rep = t.analyze(pdf, cfg, "rapidocr")
    got = items(rep)
    assert rep.engine.startswith("テキスト層")
    for name, _, v in truth:
        assert got[name].text == v, name
    for name in ("総合判定", "感じ 下限", "感じ 秤量"):  # 手書きマークは画像から判定する
        assert got[name].text == EXPECTED[name]
    assert rep.verdict == "合格"


@pytest.mark.parametrize("kind, engine", [("digital", "テキスト層 + rapidocr"), ("scan", "rapidocr")])
def test_sample_pdfs(cfg, kind, engine):
    rep = t.analyze(ROOT / "sample" / f"denshitenbin_kensa_seisekisho01_{kind}.pdf", cfg, "rapidocr")
    got = items(rep)
    assert rep.engine == engine
    for name, v in EXPECTED.items():
        assert got[name].text == v, name
    assert rep.verdict == "合格"


def test_digital_pdf_out_of_spec_is_ng(cfg, aligned, truth, tmp_path):
    texts = [(box, "+3" if name == "四隅 調整後 右後" else v) for name, box, v in truth if v]
    rep = t.analyze(_pdf(tmp_path / "ng.pdf", aligned, texts, whiten=True), cfg, "rapidocr")
    assert [it.name for it in rep.items if it.cell.status == NG] == ["四隅 調整後 右後"]
    assert rep.verdict == "不合格"


def test_invisible_ocr_text_is_ignored(cfg, aligned, truth, tmp_path):
    # スキャナーが OCR で埋め込む透明な文字（ここでは誤った値）は使わず、画像を OCR で読む
    texts = [(box, "9") for name, box, _ in truth if name in NUMBERS]
    pdf = _pdf(tmp_path / "searchable.pdf", aligned, texts, invisible=True)
    page = di.load_source(pdf)
    assert page.words is None
    got = items(t.analyze(pdf, cfg, "rapidocr"))
    for name, v in NUMBERS.items():
        assert got[name].text == v, name


def test_scanned_pdf_pages(cfg, capsys, tmp_path):
    gray = np.array(Image.open(SAMPLE).convert("L"))
    pdf = tmp_path / "scan.pdf"
    Image.fromarray(gray).save(pdf, "PDF", resolution=150, save_all=True, append_images=[Image.fromarray(gray)])
    assert di.page_count(pdf) == 2
    assert io.main([str(pdf), "--config", str(CONFIG), "--engine", "rapidocr"]) == 0
    out = capsys.readouterr().out
    assert "(1/2 ページ)" in out and "(2/2 ページ)" in out
    assert out.count("■ 総合判定: 合格") == 2


# ---------------------------------------------------------------- Excel

ROWS = ["長さ", "幅", "高さ"]


def _workbook(path, values, *, width_b=14, judgment="OK"):
    """検査成績書風の Excel。B〜D 列が計測値（3 行 × 3 回）、F2 が判定、B7 が数式。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "成績書"
    thin = Side(style="thin")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws["A1"] = "部品検査成績書"
    ws["A1"].font = Font(sz=16, bold=True)
    ws.merge_cells("A1:D1")
    ws["E2"], ws["F2"] = "判定", judgment
    for j, head in enumerate(["項目", "1回目", "2回目", "3回目"]):
        ws.cell(2, 1 + j, head)
    for i, (name, vals) in enumerate(zip(ROWS, values)):
        ws.cell(3 + i, 1, name)
        for j, v in enumerate(vals):
            c = ws.cell(3 + i, 2 + j, v)
            c.number_format = "0.00"
            c.alignment = Alignment(horizontal="center")
    ws["A7"], ws["B7"] = "合計", "=SUM(B3:D5)"  # openpyxl で保存した数式は計算結果が無い
    for row in ws["A2:F5"]:
        for c in row:
            c.border = box
    ws.column_dimensions["A"].width = 10
    for col in "BCD":
        ws.column_dimensions[col].width = width_b
    wb.save(path)
    return path


def _excel_config(tmp_path, ref):
    boxes = di.cell_boxes(ref)
    x0, y0 = boxes["B3"][:2]
    x1, y1 = boxes["D5"][2:]
    return {
        "type": "template", "name": "Excel 様式", "reference_image": ref.name, "engine": "rapidocr",
        "min_confidence": 60, "_dir": str(tmp_path),
        "fields": [
            {"name": "計測値", "kind": "table", "box": [x0, y0, x1, y1], "rows": 3, "cols": 3,
             "row_labels": ROWS, "min": 9.9, "max": 10.1, "decimals": 2},
            {"name": "判定", "kind": "text", "box": list(boxes["F2"]), "expect": ["OK"], "ng_values": ["NG"]},
            {"name": "合計", "kind": "number", "box": list(boxes["B7"])},
        ],
    }


GOOD = [[10, 10.02, 9.98], [10.01, 10, 10], [9.99, 10.05, 10]]


def test_excel_template(tmp_path):
    ref = _workbook(tmp_path / "ref.xlsx", GOOD)
    cfg = _excel_config(tmp_path, ref)
    assert t.detect_grid(t.load_page(ref), cfg["fields"][0]["box"]) == (3, 3)  # 罫線から表を検出できる

    bad = copy.deepcopy(GOOD)
    bad[1][2] = 10.2  # 規格外
    path = _workbook(tmp_path / "in.xlsx", bad, width_b=16, judgment="NG")  # 列幅が違っても位置合わせする
    rep = t.analyze(path, cfg, "rapidocr")
    got = items(rep)
    assert rep.engine.startswith("テキスト層")
    assert got["計測値 長さ 1"].text == "10.00"  # 表示形式（0.00）どおりに読む
    assert got["計測値 幅 3"].status == NG
    assert (got["判定"].text, got["判定"].status) == ("NG", NG)
    assert got["合計"].status == CHECK  # 数式の計算結果が保存されていない → 空欄として要確認
    assert sum(it.cell.status == OK for it in rep.items) == 8
    assert rep.verdict == "要目視確認"


@pytest.mark.parametrize(
    "value, fmt, text",
    [(1, "0.000", "1.000"), (0.1, "0.00", "0.10"), (4, "+0;-0;0", "+4"), (-1, "+0;-0;0", "-1"),
     (0, "+0;-0;0", "0"), (100.0003, "General", "100.0003"), (10.0, "General", "10"),
     (1234.5, "#,##0.0", "1,234.5"), (0.256, "0.0%", "25.6%"), (-0.004, "0.00", "0.00"),
     (2.5, "0", "3"), (3, '0" g"', "3 g"), (12, "@", "12"), (0.5, "0.0#", "0.5"), (0.555, "0.0#", "0.56"),
     (-2, "0.00;[Red]-0.00", "-2.00"), ("OK", "General", "OK"), (None, "0.00", "")],
)
def test_excel_number_format(value, fmt, text):
    assert di.format_value(value, fmt) == text
